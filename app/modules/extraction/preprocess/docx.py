"""DOCX preparation.

A DOCX file is a ZIP archive of XML files. Before opening it with python-docx we
inspect the ZIP itself, because a malicious archive can be built to explode
into gigabytes when unpacked (a "zip bomb"), or can carry macros.

The document is converted to Markdown-like text (paragraphs, then tables), which
models read well. Embedded images are not extracted in release 1.
"""

from __future__ import annotations

import io
import zipfile

import docx

from app.core.errors import JobError

_MAX_ENTRIES = 2000
_MAX_TOTAL_UNCOMPRESSED = 50 * 1024 * 1024    # 50 MB
_MAX_RATIO = 100                              # uncompressed / compressed, per entry
_CHARS_PER_PAGE = 3000                        # how we count "pages" for a DOCX (FR-2.5, usage)


def check_zip(data: bytes) -> None:
    """Reject zip bombs and macro-enabled documents (FR-2.4) before parsing."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise JobError("corrupt_file") from None
    infos = zf.infolist()
    if len(infos) > _MAX_ENTRIES or sum(i.file_size for i in infos) > _MAX_TOTAL_UNCOMPRESSED:
        raise JobError("file_too_complex")
    if any(i.compress_size and i.file_size / i.compress_size > _MAX_RATIO for i in infos):
        raise JobError("file_too_complex")
    if any(i.filename.lower().endswith("vbaproject.bin") for i in infos):
        raise JobError("macro_not_allowed")


def docx_to_text(data: bytes) -> str:
    """A DOCX as plain text: paragraphs first, then each table as ' | '-separated rows."""
    check_zip(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise JobError("corrupt_file") from None
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = [[cell.text.strip().replace("|", "/") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        # Markdown table: header row, separator row, body rows.
        parts.append("| " + " | ".join(rows[0]) + " |")
        parts.append("|" + "---|" * len(rows[0]))
        parts.extend("| " + " | ".join(r) + " |" for r in rows[1:])
    return "\n".join(parts)


def docx_pages(data: bytes, max_pages: int) -> list[dict]:
    """Split the DOCX text into page-sized chunks. DOCX pages have no image."""
    text = docx_to_text(data)
    chunks = [text[i:i + _CHARS_PER_PAGE] for i in range(0, max(len(text), 1), _CHARS_PER_PAGE)]
    if len(chunks) > max_pages:
        raise JobError("too_many_pages")
    return [{"text": chunk, "image": None, "scanned": False} for chunk in chunks]
